"""Build an evidence register and chain-of-custody log for an intake folder.

Hashes every file at intake so that any later claim about integrity is
provable. Run this BEFORE opening anything substantive. It cannot be done
retroactively.

Usage:
    python evidence_register.py --input <folder> --output <register.xlsx>
        [--matter "Acme Holdings 2025 books review"]
        [--received-from "the bookkeeper"]
        [--received-how "Email attachment"]

Re-verify a previously built register against the files as they stand now:
    python evidence_register.py --verify <register.xlsx> --input <folder>

The verify mode answers one question: is what I am analyzing still what I was
given? It reports UNCHANGED, CHANGED, MISSING, and NEW per item.
"""

import argparse
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from audit_common import (  # noqa: E402
    BAD_FILL, GOOD_FILL, new_workbook, provenance_banner, read_table,
    save_workbook, sha256_file, stamp, write_sheet,
)

SKIP_NAMES = {".DS_Store", "Thumbs.db", "desktop.ini"}
SKIP_DIRS = {".git", "__pycache__", ".ipynb_checkpoints"}

DESCRIBE = [
    (("bank", "stmt", "statement"), "Bank or card statement"),
    (("gl", "general ledger", "ledger"), "General ledger"),
    (("trial", "tb"), "Trial balance"),
    (("p&l", "pl", "profit", "income"), "Profit and loss"),
    (("balance sheet", "bs"), "Balance sheet"),
    (("1099", "w-2", "w2", "k-1", "k1"), "Tax information return"),
    (("return", "1040", "1065", "1120"), "Tax return"),
    (("invoice", "inv"), "Invoice"),
    (("payroll",), "Payroll record"),
    (("agreement", "contract", "operating"), "Agreement or contract"),
    (("minutes", "consent", "resolution"), "Governance record"),
    (("recon",), "Reconciliation"),
    (("journal", "je"), "Journal entries"),
    (("audit trail", "changelog", "change log"), "System audit trail"),
]


def guess_description(name):
    low = name.lower()
    for keys, label in DESCRIBE:
        for key in keys:
            if key in low:
                return label
    return ""


def walk_files(root):
    found = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
        for name in sorted(filenames):
            if name in SKIP_NAMES or name.endswith(".tmp"):
                continue
            full = os.path.join(dirpath, name)
            if os.path.islink(full):
                continue
            found.append(full)
    return sorted(found)


def collect(root, received_from, received_how):
    rows = []
    for idx, full in enumerate(walk_files(root), start=1):
        rel = os.path.relpath(full, root)
        st = os.stat(full)
        rows.append({
            "id": "E-%03d" % idx,
            "rel": rel.replace("\\", "/"),
            "name": os.path.basename(full),
            "ext": os.path.splitext(full)[1].lower().lstrip("."),
            "size": st.st_size,
            "sha256": sha256_file(full),
            "modified": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            "logged": stamp(),
            "from": received_from,
            "how": received_how,
            "desc": guess_description(os.path.basename(full)),
            "abs": full,
        })
    return rows


def build_register(args):
    root = os.path.abspath(args.input)
    if not os.path.isdir(root):
        raise SystemExit("Input folder not found: %s" % root)

    items = collect(root, args.received_from, args.received_how)
    if not items:
        raise SystemExit("No files found under %s" % root)

    wb = new_workbook()
    banner = provenance_banner(
        "evidence_register.py",
        "matter=%s; root=%s" % (args.matter or "(unnamed)", root),
    )

    headers = [
        "Item ID", "File name", "Path as received", "Type", "Size (bytes)",
        "SHA-256", "File modified", "Logged at", "Received from",
        "How received", "Description", "Period covered", "Custody location",
        "Notes",
    ]
    rows = [[
        it["id"], it["name"], it["rel"], it["ext"], it["size"], it["sha256"],
        it["modified"], it["logged"], it["from"], it["how"], it["desc"],
        "", root, "",
    ] for it in items]

    write_sheet(
        wb, "Evidence Register", headers, rows,
        widths=[10, 34, 40, 8, 14, 66, 20, 20, 24, 18, 26, 18, 40, 30],
        notes=[
            banner,
            "Period covered and Notes are filled in by the examiner. The "
            "period a document COVERS is not the date it was received.",
        ],
    )

    custody_headers = [
        "Date and time", "Item ID", "Action", "By whom", "Purpose",
        "Hash at this point", "Notes",
    ]
    custody_rows = [[
        it["logged"], it["id"], "Received and hashed at intake",
        args.examiner, "Intake and preservation", it["sha256"], "",
    ] for it in items]
    write_sheet(
        wb, "Chain of Custody", custody_headers, custody_rows,
        widths=[20, 10, 34, 24, 30, 66, 34],
        notes=[
            "Append a row for EVERY transfer, copy, conversion, or analysis "
            "step. An entry made later is worth far less than one made at "
            "the time.",
        ],
    )

    write_sheet(
        wb, "Transformation Log",
        ["Date and time", "Item ID", "Step", "Tool and version",
         "Parameters and thresholds", "Output produced", "By whom", "Notes"],
        [],
        widths=[20, 10, 34, 24, 40, 34, 20, 34],
        notes=[
            "Log every change of shape: extraction, export, normalization, "
            "merge, dedupe, filter. Exclusions count, including ones added "
            "only to reduce noise.",
        ],
    )

    write_sheet(
        wb, "Requested Not Produced",
        ["Item requested", "Date requested", "Requested from", "Status",
         "What it would have let me test", "Effect on conclusions"],
        [],
        widths=[40, 16, 24, 16, 44, 40],
        notes=[
            "Anything still open at report time goes into Scope and "
            "Limitations BY NAME. An unrecorded gap becomes an unexplained "
            "hole in the analysis later.",
        ],
    )

    summary = [
        ["Matter", args.matter or "(unnamed)"],
        ["Intake root", root],
        ["Files logged", len(items)],
        ["Total bytes", sum(i["size"] for i in items)],
        ["Register built", stamp()],
        ["Examiner", args.examiner],
        ["Hash algorithm", "SHA-256 over raw bytes"],
    ]
    write_sheet(wb, "Summary", ["Field", "Value"], summary, widths=[24, 80])

    out = save_workbook(wb, args.output)
    print("Evidence register written: %s" % out)
    print("Files logged: %d" % len(items))
    print("Total bytes: %d" % sum(i["size"] for i in items))
    print("")
    print("NEXT: set the intake folder read-only and work only on copies.")
    return 0


def verify(args):
    reg_path = os.path.abspath(args.verify)
    root = os.path.abspath(args.input)
    if not os.path.isfile(reg_path):
        raise SystemExit("Register not found: %s" % reg_path)
    if not os.path.isdir(root):
        raise SystemExit("Input folder not found: %s" % root)

    from openpyxl import load_workbook
    probe = load_workbook(reg_path, read_only=True)
    sheets = list(probe.sheetnames)
    probe.close()
    if "Evidence Register" not in sheets:
        raise SystemExit(
            "%s has no 'Evidence Register' sheet. Sheets present: %s\n"
            "This does not look like a register produced by this script. "
            "Nothing was compared."
            % (reg_path, ", ".join(sheets)))

    headers, recorded = read_table(
        reg_path, sheet="Evidence Register",
        header_contains=["Item ID", "Path as received", "SHA-256"])
    by_rel = {}
    for row in recorded:
        rel = str(row.get("Path as received") or "").strip()
        if rel:
            by_rel[rel] = row

    # Positive control on the instrument itself. If the register parsed to
    # nothing, every current file reads as NEW and the run reports no changes
    # and no missing items, which is indistinguishable from a clean result.
    # A verification that compared nothing must never exit as though it passed.
    if not by_rel:
        raise SystemExit(
            "Read 0 usable rows from the Evidence Register sheet of %s.\n"
            "Nothing was compared. This is an instrument failure, not a clean "
            "result. Check that the file is a register produced by this script."
            % reg_path)

    current = {}
    for full in walk_files(root):
        rel = os.path.relpath(full, root).replace("\\", "/")
        current[rel] = full

    results = []
    changed = missing = new = unchanged = 0

    for rel, row in sorted(by_rel.items()):
        item_id = str(row.get("Item ID") or "")
        want = str(row.get("SHA-256") or "").strip().lower()
        if rel not in current:
            results.append([item_id, rel, "MISSING", want, "", "File is no longer present"])
            missing += 1
            continue
        got = sha256_file(current[rel]).lower()
        if got == want:
            results.append([item_id, rel, "UNCHANGED", want, got, ""])
            unchanged += 1
        else:
            results.append([item_id, rel, "CHANGED", want, got,
                            "Content differs from intake. Do not analyze until explained."])
            changed += 1

    for rel in sorted(current):
        if rel not in by_rel:
            results.append(["", rel, "NEW", "", sha256_file(current[rel]),
                            "Present now, not in the register. Log it."])
            new += 1

    wb = new_workbook()
    ws = write_sheet(
        wb, "Verification",
        ["Item ID", "Path", "Status", "Hash at intake", "Hash now", "Note"],
        results,
        widths=[10, 44, 14, 66, 66, 46],
        notes=[provenance_banner("evidence_register.py --verify",
                                 "register=%s; root=%s" % (reg_path, root))],
    )

    first_data_row = ws.max_row - len(results) + 1
    for offset, row in enumerate(results):
        excel_row = first_data_row + offset
        fill = GOOD_FILL if row[2] == "UNCHANGED" else BAD_FILL
        for col in range(1, 7):
            ws.cell(row=excel_row, column=col).fill = fill

    out = save_workbook(wb, args.output or (os.path.splitext(reg_path)[0] + " - verification.xlsx"))
    print("Verification written: %s" % out)
    print("UNCHANGED %d | CHANGED %d | MISSING %d | NEW %d" % (unchanged, changed, missing, new))
    if changed or missing:
        print("")
        print("STOP. At least one item is not what it was at intake.")
        print("Do not report findings from a changed or missing source until it is explained.")
        return 2
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True, help="Intake folder to hash")
    ap.add_argument("--output", help="Output .xlsx path")
    ap.add_argument("--matter", default="", help="Matter or engagement name")
    ap.add_argument("--received-from", default="", help="Who produced these records")
    ap.add_argument("--received-how", default="", help="Email, portal, drive, direct export")
    ap.add_argument("--examiner", default=os.environ.get("USERNAME", "examiner"))
    ap.add_argument("--verify", help="Re-verify against an existing register")
    args = ap.parse_args()

    if args.verify:
        return verify(args)
    if not args.output:
        ap.error("--output is required when building a register")
    return build_register(args)


if __name__ == "__main__":
    raise SystemExit(main())
